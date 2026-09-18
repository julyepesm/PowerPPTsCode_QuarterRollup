"""
Slide Templates - Refactored Modular Version

This module orchestrates slide creation by coordinating:
- tactic_config: Business rules and tactic definitions
- slide_layouts: Slide structure and positioning
- slide_components: Data visualizations

The SlideTemplates class now acts as a high-level orchestrator rather than
handling everything itself.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


# Import our modular components
from tactic_config import TacticConfig
from slide_layouts import SlideLayouts
from slide_components import SlideComponents


class SlideTemplates:
    """Orchestrates PowerPoint slide creation using modular components"""

    def __init__(self, brand_configs):
        """
        Initialize with brand configurations and set up modular components
        
        Args:
            brand_configs: BrandConfigurations instance with brand-specific settings
        """
        self.brand_configs = brand_configs
        
        # Initialize modular components
        self.tactic_config = TacticConfig()
        self.layouts = SlideLayouts(brand_configs)
        self.components = SlideComponents(brand_configs, self.tactic_config, base_path=brand_configs.base_path)
    
    # ========== DELEGATION METHODS (for backward compatibility) ==========
    
    def get_tactic_metrics(self, tactic, strategy=None):
        """Get metrics for a tactic - delegates to TacticConfig"""
        return self.tactic_config.get_tactic_metrics(tactic, strategy)
    
    def get_tactic_benchmarks(self, tactic, strategy=None):
        """Get benchmarks for a tactic - delegates to TacticConfig"""
        return self.tactic_config.get_tactic_benchmarks(tactic, strategy)
    
    def get_tactic_chart_config(self, tactic, strategy=None):
        """Get chart config for a tactic - delegates to TacticConfig"""
        return self.tactic_config.get_tactic_chart_config(tactic, strategy)
    
    def pixels_to_inches(self, pixels):
        """Convert pixels to inches - delegates to layouts module"""
        return self.layouts.pixels_to_inches(pixels)
    
    def format_number(self, value):
        """Format numbers for display - delegates to components module"""
        return self.components.format_number(value)
    
    def format_percentage(self, value):
        """Format percentages - delegates to components module"""
        return self.components.format_percentage(value)
    
    # ========== HIGH-LEVEL SLIDE CREATION METHODS ==========
    
    def create_title_slide(self, brand_name, market_name, month_year, zone_lma_type=None):
        """
        Create title slide using modular layout components
        
        Args:
            brand_name: Name of the brand
            market_name: Market/LMA name
            month_year: Month and year string
            zone_lma_type: Optional zone/LMA type indicator
        
        Returns:
            tuple: (Presentation object, slide object)
        """
        # Create blank presentation with standard widescreen dimensions
        prs = Presentation()
        prs.slide_width = Inches(13.33)   # 16:9
        prs.slide_height = Inches(7.5)
        
        # Remove default slide if it exists
        if prs.slides:
            prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])
        
        # Add blank slide
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        
        # Use layouts module to build the complete title slide
        self.layouts.build_title_slide(
            slide=slide,
            brand_name=brand_name,
            market_name=market_name,
            month_year=month_year,
            zone_lma_type=zone_lma_type
        )
        
        print(f"Created title slide for {brand_name} - {market_name}")
        return prs, slide
    
    def create_data_slide(self, brand_name, market_name, month_year, 
                         current_metrics=None, historical_data=None, 
                         tactic=None, audience_label=None, zone_lma_type=None, 
                         strategy=None, existing_prs=None):
        """
        Create data slide with Power BI metrics
        
        Args:
            brand_name: Name of the brand
            market_name: Market/LMA name
            month_year: Month and year string
            current_metrics: DataFrame with current period metrics
            historical_data: DataFrame with historical metrics
            tactic: Tactic name
            audience_label: Audience segment label
            zone_lma_type: Optional zone/LMA type indicator
            strategy: Optional strategy for Display tactic splitting
            existing_prs: Optional existing presentation to add to
        
        Returns:
            tuple: (Presentation object, slide object)
        """
        # Note: Metrics are already calculated in main_slide_generator.py
        
        # Use existing presentation or create new one
        if existing_prs is None:
            prs = Presentation()
            prs.slide_width = Inches(13.33)
            prs.slide_height = Inches(7.5)
            
            if prs.slides:
                prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])
        else:
            prs = existing_prs
        
        # Add blank slide
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        
        # Resolve Display tactic based on strategy
        resolved_tactic = tactic
        if tactic == 'Display' and strategy:
            try:
                from display_resolution import resolve_display_tactic
                resolved_tactic = resolve_display_tactic(tactic, strategy)
                print(f"  Resolved {tactic} + {strategy} -> {resolved_tactic} for colors")
            except ImportError:
                resolved_tactic = 'BT Display'  # Default fallback
        elif tactic == 'Display':
            resolved_tactic = 'BT Display'  # Default when no strategy

        # Get brand configuration
        brand_config = self.brand_configs.get_brand_config(brand_name, tactic=resolved_tactic)
        
        # Build slide using modular components
        self._assemble_data_slide(
            slide=slide,
            brand_config=brand_config,
            market_name=market_name,
            month_year=month_year,
            current_metrics=current_metrics,
            historical_data=historical_data,
            tactic=tactic,
            audience_label=audience_label,
            zone_lma_type=zone_lma_type,
            strategy=strategy
        )
        
        print(f"Created data slide for {brand_name} - {market_name} - {tactic}")
        return prs, slide
    
    def create_combined_presentation(self, brand_name, market_name, month_year, 
                                    current_metrics=None, historical_data=None, 
                                    tactic=None, audience_label=None, 
                                    zone_lma_type=None, strategy=None):
        """
        Create a presentation with both title slide (slide 1) and data slide (slide 2)
        
        Args:
            brand_name: Name of the brand
            market_name: Market/LMA name
            month_year: Month and year string
            current_metrics: DataFrame with current period metrics
            historical_data: DataFrame with historical metrics
            tactic: Tactic name
            audience_label: Audience segment label
            zone_lma_type: Optional zone/LMA type indicator
            strategy: Optional strategy for Display tactic splitting
        
        Returns:
            Presentation object with both slides
        """
        # Note: Metrics are already calculated in main_slide_generator.py
        
        # Create presentation
        prs = Presentation()
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)
        
        if prs.slides:
            prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])
        
        blank_slide_layout = prs.slide_layouts[6]
        brand_config = self.brand_configs.get_brand_config(brand_name)
        
        # Slide 1: Title slide
        title_slide = prs.slides.add_slide(blank_slide_layout)
        self.layouts.build_title_slide(
            slide=title_slide,
            brand_name=brand_name,
            market_name=market_name,
            month_year=month_year,
            zone_lma_type=zone_lma_type
        )
        
        # Slide 2: Data slide
        data_slide = prs.slides.add_slide(blank_slide_layout)
        self._assemble_data_slide(
            slide=data_slide,
            brand_config=brand_config,
            market_name=market_name,
            month_year=month_year,
            current_metrics=current_metrics,
            historical_data=historical_data,
            tactic=tactic,
            audience_label=audience_label,
            zone_lma_type=zone_lma_type,
            strategy=strategy
        )
        
        print(f"Created combined presentation: Title + Data slide for {brand_name} - {market_name}")
        return prs
    
    # ========== INTERNAL ASSEMBLY METHOD ==========
    
    def _assemble_data_slide(self, slide, brand_config, market_name, month_year,
                            current_metrics, historical_data, tactic, 
                            audience_label, zone_lma_type, strategy=None):
        """
        Assemble a complete data slide using modular components
        
        This method coordinates the layouts and components modules to build
        a complete data slide with structure and visualizations.
        
        Args:
            slide: PowerPoint slide object
            brand_config: Brand configuration dictionary
            market_name: Market/LMA name
            month_year: Month and year string
            current_metrics: DataFrame with current period metrics
            historical_data: DataFrame with historical metrics
            tactic: Tactic name (may be "Display")
            audience_label: Audience segment label
            zone_lma_type: Optional zone/LMA type indicator
            strategy: Optional strategy for Display tactic splitting
        """
        # Extract brand name from brand_config
        brand = brand_config['name']  # e.g., "GMC", "Buick", "Chevrolet", "Cadillac"
        
        # ========== EXTRACT AUDIENCE FROM AUDIENCE_LABEL ==========
        # audience_label formats: "Meta Display - GEN", "YouTube - HIS", "BT Display - ASN"
        audience = None
        if audience_label:
            # Split by " - " and take the last part
            parts = audience_label.split(' - ')
            if len(parts) > 1:
                audience = parts[-1].strip()  # Gets "GEN", "HIS", "ASN"
            else:
                # Fallback: assume the whole thing is the audience
                audience = audience_label.strip()
        
        # If we still don't have audience, try to extract from current_metrics
        if not audience and current_metrics is not None and not current_metrics.empty:
            if 'PoP Master Table[Audience]' in current_metrics.columns:
                audience = current_metrics['PoP Master Table[Audience]'].iloc[0]
            elif '[Audience]' in current_metrics.columns:
                audience = current_metrics['[Audience]'].iloc[0]
        
        # Default to 'GEN' if we still can't determine audience
        if not audience:
            audience = 'GEN'
            print(f"  [WARN] Could not determine audience, defaulting to 'GEN'")
        
        print(f"  Extracted audience: {audience}")
        # ========== END AUDIENCE EXTRACTION ==========
        
        # Resolve Display tactic for slide title
        # This ensures slide title shows "BT Display" or "Retargeting Display"
        resolved_tactic = self.tactic_config._resolve_display_tactic(tactic, strategy)
        
        # Step 1: Build slide structure using layouts module
        # Use resolved_tactic for the slide title
        self.layouts.build_data_slide_structure(
            slide=slide,
            brand_name=brand_config['name'],
            market_name=market_name,
            month_year=month_year,
            tactic=resolved_tactic,  # Pass resolved tactic
            audience_label=audience_label,
            zone_lma_type=zone_lma_type
        )
        
        # Step 2: Add data visualizations using components module
        
        # Add metric boxes if current data available
        if current_metrics is not None and not current_metrics.empty:
            self.components.add_metric_boxes(
                slide=slide,
                current_metrics=current_metrics,
                brand_config=brand_config,
                tactic=tactic,
                strategy=strategy,
                brand=brand,
                audience=audience
            )
        
        # Add benchmark labels
        self.components.add_benchmark_labels(
            slide=slide,
            brand_config=brand_config,
            tactic=tactic,
            strategy=strategy,
            brand=brand,
            audience=audience
        )
        
        # Add KBA table and chart if historical data available
        if historical_data is not None and not historical_data.empty:
            self.components.add_kba_table(
                slide=slide,
                historical_data=historical_data,
                brand_config=brand_config,
                month_year=month_year
            )
            
            # Add combination chart with benchmark line
            # Add combination chart with benchmark line
            self.components.add_tactic_chart(
                slide=slide,
                historical_data=historical_data,
                brand_config=brand_config,
                brand=brand,           # Now properly defined
                tactic=tactic,         # Now properly defined
                audience=audience,     # Now properly defined (extracted above)
                month_year=month_year,
                strategy=strategy,      # Now properly defined
                market_name=market_name # Pass market name for error logging
            )
        
        print(f"Assembled complete data slide for {tactic}")
    
    # ========== LEGACY PRIVATE METHODS (kept for compatibility) ==========
    # These can be removed once all calling code is updated
    
    def _add_title_background(self, slide, brand_name, brand_config):
        """Legacy method - delegates to layouts module"""
        self.layouts.add_title_background(slide, brand_name, brand_config)
    
    def _add_title_brand_logo(self, slide, brand_config):
        """Legacy method - delegates to layouts module"""
        self.layouts.add_title_brand_logo(slide, brand_config)
    
    def _add_mrg_logo(self, slide):
        """Legacy method - delegates to layouts module"""
        self.layouts.add_mrg_logo(slide, position='title')
    
    def _add_gm_confidential(self, slide, brand_config):
        """Legacy method - delegates to layouts module"""
        self.layouts.add_gm_confidential(slide, brand_config, position='title')
    
    def _add_title_text(self, slide, brand_config, market_name, zone_lma_type):
        """Legacy method - delegates to layouts module"""
        self.layouts.add_title_text(slide, brand_config, market_name, zone_lma_type)
    
    def _add_title_line(self, slide, brand_config):
        """Legacy method - delegates to layouts module"""
        self.layouts.add_title_line(slide, brand_config)
    
    def _add_market_info(self, slide, brand_config, market_name, month_year):
        """Legacy method - delegates to layouts module"""
        self.layouts.add_market_info(slide, brand_config, market_name, month_year)
    
    def _add_data_background(self, slide, brand_config):
        """Legacy method - delegates to layouts module"""
        self.layouts.add_data_background(slide, brand_config)
    
    def _add_data_brand_logo(self, slide, brand_config):
        """Legacy method - delegates to layouts module"""
        self.layouts.add_data_brand_logo(slide, brand_config)

    def create_error_summary_slide(self, prs, error_list):
        """
        Create a slide summarizing data validation errors.
        
        Args:
            prs: PowerPoint presentation object
            error_list: List of error dictionaries from ErrorCollector
        """
        if not error_list:
            return
            
        # Create a blank slide
        slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blind layout
        
        # Add Title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12.33), Inches(1))
        title_frame = title_box.text_frame
        title_p = title_frame.paragraphs[0]
        title_p.text = "Data Anomalies & Validation Errors"
        title_p.font.size = Pt(24)
        title_p.font.bold = True
        
        # Add Table
        rows = len(error_list) + 1
        cols = 5 # Market, Brand, Tactic, Month, Issue
        
        table_width = Inches(12)
        table_height = Inches(0.4 * rows)
        
        # Guard against too many errors fitting on one slide
        # For now, just truncating if extremely large, or letting it flow off?
        # Let's cap at 15 errors to avoid breaking slide, maybe? 
        # Or just let it spill.
        
        left = Inches(0.6)
        top = Inches(1.5)
        
        shape = slide.shapes.add_table(rows, cols, left, top, table_width, table_height)
        table = shape.table
        
        # Headers
        headers = ["Market", "Brand", "Tactic", "Month", "Anomaly"]
        for i, h in enumerate(headers):
            cell = table.cell(0, i)
            cell.text = h
            cell.text_frame.paragraphs[0].font.bold = True
            cell.text_frame.paragraphs[0].font.size = Pt(11)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(200, 200, 200)
            
        # Data
        for i, err in enumerate(error_list):
            row_idx = i + 1
            table.cell(row_idx, 0).text = str(err.get('market', ''))
            table.cell(row_idx, 1).text = str(err.get('brand', ''))
            table.cell(row_idx, 2).text = str(err.get('tactic', ''))
            table.cell(row_idx, 3).text = str(err.get('month', ''))
            table.cell(row_idx, 4).text = str(err.get('message', ''))
            
            # Formatting
            for c in range(5):
                table.cell(row_idx, c).text_frame.paragraphs[0].font.size = Pt(10)
        
        print(f"  [INFO] Created error summary slide with {len(error_list)} entries.")
    
    def _add_data_logo_divider(self, slide):
        """Legacy method - delegates to layouts module"""
        self.layouts.add_data_logo_divider(slide, brand_config=None)
    
    def _add_data_slide_title(self, slide, market_name, subtitle, brand_config):
        """Legacy method - delegates to layouts module"""
        self.layouts.add_data_slide_title(slide, market_name, subtitle, brand_config)
    
    def _add_audience_label(self, slide, audience_label, brand_config):
        """Legacy method - delegates to layouts module"""
        self.layouts.add_audience_label(slide, audience_label, brand_config)
    
    def _add_source_label(self, slide, brand_config, zone_lma_type, month_year):
        """Legacy method - delegates to layouts module"""
        self.layouts.add_source_label(slide, brand_config, zone_lma_type, month_year)
    
    def _add_top_metrics(self, slide, current_metrics, brand_config, tactic, strategy=None):
        """Legacy method - delegates to components module"""
        self.components.add_metric_boxes(slide, current_metrics, brand_config, tactic, strategy)
    
    def _add_benchmark_labels(self, slide, brand_config, tactic):
        """Legacy method - delegates to components module"""
        self.components.add_benchmark_labels(slide, brand_config, tactic)
    
    def _add_kba_table(self, slide, historical_data, brand_config, month_year):
        """Legacy method - delegates to components module"""
        self.components.add_kba_table(slide, historical_data, brand_config, month_year)
    
    def _add_tactic_chart(self, slide, historical_data, brand_config, tactic, month_year):
        """Legacy method - delegates to components module"""
        # This legacy method can't properly support the new chart - would need brand and audience
        print("[WARN] Warning: Using legacy _add_tactic_chart method - chart may not work correctly")
        self.components.add_tactic_chart(slide, historical_data, brand_config, 'GMC', tactic, 'GEN', month_year, None)
    
    def _add_data_layout(self, slide, brand_config, market_name, month_year, 
                        current_metrics, historical_data, tactic, audience_label, 
                        zone_lma_type, strategy=None):
        """Legacy method - now calls _assemble_data_slide"""
        self._assemble_data_slide(
            slide, brand_config, market_name, month_year,
            current_metrics, historical_data, tactic, 
            audience_label, zone_lma_type, strategy
        )


# ========== USAGE EXAMPLE ==========

if __name__ == "__main__":
    """
    Example usage of the refactored SlideTemplates class
    """
    
    # Assuming you have your brand_configs set up
    # from brand_configurations import BrandConfigurations
    # brand_configs = BrandConfigurations()
    
    # Initialize the refactored slide templates
    # slide_templates = SlideTemplates(brand_configs)
    
    # Create a title slide
    # prs, title_slide = slide_templates.create_title_slide(
    #     brand_name='Chevrolet',
    #     market_name='Chicago',
    #     month_year='October 2024',
    #     zone_lma_type='Zone'
    # )
    
    # Add a data slide to the same presentation
    # prs, data_slide = slide_templates.create_data_slide(
    #     brand_name='Chevrolet',
    #     market_name='Chicago',
    #     month_year='October 2024',
    #     current_metrics=current_df,
    #     historical_data=historical_df,
    #     tactic='YouTube',
    #     audience_label='YouTube - GEN',
    #     zone_lma_type='Zone',
    #     existing_prs=prs  # Add to existing presentation
    # )
    
    # prs.save('output.pptx')
    
    print("Refactored SlideTemplates ready to use!")
    print("\nKey improvements:")
    print("[OK] Modular architecture with clear separation of concerns")
    print("[OK] Tactic logic separated from presentation code")
    print("[OK] Layout and components are independent modules")
    print("[OK] Easy to test, maintain, and extend")
    print("[OK] Backward compatible with existing code")
    print("[OK] Proper audience extraction for benchmark charts")