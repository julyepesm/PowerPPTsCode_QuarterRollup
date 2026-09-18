import os
import sys
import tempfile
import unittest
from io import BytesIO

import pandas as pd
from pptx import Presentation
from pptx.util import Inches


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_PATH = os.path.dirname(SCRIPT_DIR)
os.environ.setdefault(
    "MPLCONFIGDIR", os.path.join(tempfile.gettempdir(), "powerppts-matplotlib-test")
)
sys.path.append(SCRIPT_DIR)

from brand_configurations import BrandConfigurations
from data_queries import DataQueries
from slide_templates_refactored import SlideTemplates
from tactic_config import TacticConfig
from vehicle_catalog import get_vehicle_catalog
from vehicle_insights_summary import (
    create_vehicle_insights_slide,
    create_vehicle_summary_slide,
    prepare_vehicle_current_data,
    prepare_vehicle_historical_data,
    resolve_vehicle_image_path,
)


def metric_row(vehicle, cost, site="Site A", month=None, tactic="Search"):
    row = {
        "PoP Master Table[Vehicle]": vehicle,
        "PoP Master Table[Tactic (Reporting)]": tactic,
        "PoP Master Table[Strategy (groups)]": "",
        "PoP Master Table[Site (Reporting)]": site,
        "PoP Master Table[Audience]": "GEN",
        "[Total Cost]": cost,
        "[Impressions]": 1000,
        "[Video Completions]": 0,
        "[Video Plays]": 0,
        "[Clicks]": 100,
        "[Total Conversions]": 10,
    }
    if month is not None:
        row["Master Date Table[Month Year]"] = month
    return row


class VehicleDataPreparationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tactic_config = TacticConfig()
        cls.vehicle_catalog = get_vehicle_catalog(BASE_PATH)

    def test_current_threshold_is_inclusive_and_applied_after_combo_aggregation(self):
        rows = pd.DataFrame([
            metric_row("Enclave", 25),
            metric_row("Enclave", 25),
            metric_row("Envista", 30, site="Site A"),
            metric_row("Envista", 25, site="Site B"),
            metric_row("Envision", 49.99),
        ])

        active, excluded = prepare_vehicle_current_data(
            rows, self.tactic_config, self.vehicle_catalog, "Buick",
            minimum_spend=50,
        )

        self.assertEqual(list(active), ["Enclave"])
        self.assertEqual(active["Enclave"]["Total Cost"].sum(), 50)
        excluded_names = {item["vehicle"] for item in excluded}
        self.assertEqual(excluded_names, {"Envista", "Envision"})
        envista = next(item for item in excluded if item["vehicle"] == "Envista")
        self.assertEqual(envista["total_spend"], 55)
        self.assertEqual(envista["max_combination_spend"], 30)

    def test_historical_threshold_uses_same_grain(self):
        rows = pd.DataFrame([
            metric_row("Enclave", 50, month="2026-04-01"),
            metric_row("Enclave", 49.99, month="2026-05-01"),
            metric_row("Enclave", 30, month="2026-06-01"),
            metric_row("Enclave", 20, month="2026-06-01"),
        ])

        history = prepare_vehicle_historical_data(
            rows, self.tactic_config, self.vehicle_catalog, "Buick",
            ["Enclave"], 50,
            apply_spend_filter=True,
        )["Enclave"]

        self.assertEqual(
            history["Month Year"].dt.strftime("%Y-%m").tolist(),
            ["2026-04", "2026-06"],
        )

    def test_json_aliases_merge_data_before_spend_filter(self):
        rows = pd.DataFrame([
            metric_row("Sierra", 25),
            metric_row("Sierra LD", 25),
        ])

        active, excluded = prepare_vehicle_current_data(
            rows, self.tactic_config, self.vehicle_catalog, "GMC", 50
        )

        self.assertEqual(list(active), ["Sierra 1500"])
        self.assertEqual(active["Sierra 1500"]["Total Cost"].sum(), 50)
        self.assertEqual(active["Sierra 1500"]["Impressions"].sum(), 2000)
        self.assertEqual(excluded, [])

    def test_unmapped_vehicle_is_excluded_and_preserved_for_anomaly(self):
        rows = pd.DataFrame([metric_row("Unlisted Vehicle", 50)])

        active, excluded = prepare_vehicle_current_data(
            rows, self.tactic_config, self.vehicle_catalog, "GMC", 50
        )

        self.assertEqual(active, {})
        self.assertEqual(len(excluded), 1)
        self.assertEqual(excluded[0]["vehicle"], "Unlisted Vehicle")
        self.assertEqual(excluded[0]["source_values"], "Unlisted Vehicle")
        self.assertEqual(excluded[0]["exclusion_type"], "taxonomy")
        self.assertEqual(excluded[0]["classification"], "unmapped")
        self.assertEqual(excluded[0]["total_spend"], 50)

    def test_reference_aliases_group_before_spend_filter(self):
        rows = pd.DataFrame([
            metric_row("HummerEVTruck", 25),
            metric_row("HummerPickupTruck", 25),
            metric_row("Sierra2500", 50),
        ])

        active, excluded = prepare_vehicle_current_data(
            rows, self.tactic_config, self.vehicle_catalog, "GMC", 50
        )

        self.assertEqual(
            list(active), ["HUMMER EV Pickup", "Sierra 2500 HD"]
        )
        self.assertEqual(active["HUMMER EV Pickup"]["Total Cost"].sum(), 50)
        self.assertEqual(excluded, [])

    def test_ambiguous_and_rollup_values_are_excluded_by_json(self):
        rows = pd.DataFrame([
            metric_row("HummerEV", 100),
            metric_row("SierraHD", 100),
            metric_row("Truck", 100),
        ])

        active, excluded = prepare_vehicle_current_data(
            rows, self.tactic_config, self.vehicle_catalog, "GMC", 50
        )

        self.assertEqual(active, {})
        self.assertEqual(
            {item["vehicle"] for item in excluded},
            {"HUMMER EV", "Sierra HD", "Truck"},
        )
        self.assertTrue(
            all(item["exclusion_type"] == "taxonomy" for item in excluded)
        )

    def test_registered_model_without_assets_still_generates(self):
        rows = pd.DataFrame([metric_row("SavanaCargo", 50)])

        active, excluded = prepare_vehicle_current_data(
            rows, self.tactic_config, self.vehicle_catalog, "GMC", 50
        )

        self.assertEqual(list(active), ["Savana Cargo"])
        self.assertEqual(excluded, [])
        self.assertIsNone(resolve_vehicle_image_path(
            "Savana Cargo", "GMC", "exec", BASE_PATH
        ))


class VehicleAssetResolutionTests(unittest.TestCase):
    def test_user_reference_catalog_models_are_approved(self):
        reference_models = {
            "GMC": [
                "Canyon", "Sierra 1500", "Sierra 2500 HD",
                "Sierra 3500 HD", "Sierra EV", "HUMMER EV Pickup",
                "Terrain", "Acadia", "Yukon", "Yukon XL",
                "HUMMER EV SUV", "Savana Cargo", "Savana Passenger",
                "Savana Cutaway",
            ],
            "Buick": ["Envista", "Encore GX", "Envision", "Enclave"],
            "Cadillac": [
                "XT5", "Escalade", "Escalade ESV", "Escalade-V",
                "Escalade-V ESV", "OPTIQ", "OPTIQ-V", "LYRIQ",
                "LYRIQ-V", "VISTIQ", "ESCALADE IQ", "ESCALADE IQL",
                "CT4", "CT4-V", "CT4-V Blackwing", "CT5", "CT5-V",
                "CT5-V Blackwing", "CELESTIQ",
            ],
        }
        catalog = get_vehicle_catalog(BASE_PATH)

        for brand, models in reference_models.items():
            for model in models:
                with self.subTest(brand=brand, model=model):
                    resolution = catalog.resolve(brand, model)
                    self.assertTrue(resolution.mapped)
                    self.assertTrue(resolution.generate_slides)
                    self.assertEqual(resolution.status, "approved")

    def test_catalog_resolves_versions_and_nested_directories(self):
        cases = (
            ("Enclave", "Buick", "insights", "enclave_insights.png"),
            ("Sierra LD", "GMC", "exec", "sierra1500_exec_v2.png"),
            ("Hummer EV SUV", "GMC", "insights", "hummerevsuv_insights.png"),
            ("CT4", "Cadillac", "exec", "CT4_exec.png"),
        )
        for vehicle, brand, kind, expected_filename in cases:
            with self.subTest(vehicle=vehicle, brand=brand, kind=kind):
                path = resolve_vehicle_image_path(
                    vehicle, brand, kind, BASE_PATH
                )
                self.assertIsNotNone(path)
                self.assertEqual(os.path.basename(path), expected_filename)

    def test_ambiguous_or_missing_model_returns_none(self):
        self.assertIsNone(resolve_vehicle_image_path(
            "HummerEV", "GMC", "exec", BASE_PATH
        ))
        self.assertIsNone(resolve_vehicle_image_path(
            "Escalade IQ", "Cadillac", "insights", BASE_PATH
        ))

    def test_multiline_uses_brand_level_image_for_both_slide_types(self):
        cases = (
            ("BuickMultiline", "Buick", "BuickMultiline.png"),
            ("Cadillac Multinline", "Cadillac", "Cadillac Multiline Logo.png"),
            ("GMCMultiline", "GMC", "GMCMultiline.png"),
            ("GMC Multiline", "GMC", "GMCMultiline.png"),
        )
        for vehicle, brand, expected_filename in cases:
            for kind in ("insights", "exec"):
                with self.subTest(vehicle=vehicle, brand=brand, kind=kind):
                    path = resolve_vehicle_image_path(
                        vehicle, brand, kind, BASE_PATH
                    )
                    self.assertIsNotNone(path)
                    self.assertEqual(os.path.basename(path), expected_filename)

    def test_all_configured_catalog_assets_exist(self):
        catalog = get_vehicle_catalog(BASE_PATH)
        self.assertEqual(catalog.validate_configured_assets(), [])


class VehicleQueryTests(unittest.TestCase):
    def test_vehicle_queries_have_required_grain_and_three_month_window(self):
        queries = DataQueries()
        current = queries.get_vehicle_current_metrics_query(
            "GMC", "Example", "2026-06-01", market_code="12345"
        )
        historical = queries.get_vehicle_historical_metrics_query(
            "GMC", "Example", "2026-06-01", months=3,
            market_code="12345",
        )

        for required in ("[Vehicle]", "[Tactic (Reporting)]", "[Site (Reporting)]", "[Audience]"):
            self.assertIn(required, current)
            self.assertIn(required, historical)
        self.assertIn("-(3-1)", historical)


class VehicleSlideSmokeTests(unittest.TestCase):
    def test_creates_cadillac_pair_with_dark_theme(self):
        brand_configs = BrandConfigurations(BASE_PATH)
        slide_templates = SlideTemplates(brand_configs)
        presentation = Presentation()
        presentation.slide_width = Inches(13.33)
        presentation.slide_height = Inches(7.5)
        if presentation.slides:
            presentation.slides._sldIdLst.remove(
                presentation.slides._sldIdLst[0]
            )

        current = pd.DataFrame([
            metric_row("CT4", 100, tactic="Search"),
            metric_row("CT4", 150, site="Site B", tactic="FEP"),
        ])
        vehicle_catalog = get_vehicle_catalog(BASE_PATH)
        active, _ = prepare_vehicle_current_data(
            current, slide_templates.tactic_config, vehicle_catalog,
            "Cadillac", 50,
        )
        history_raw = pd.DataFrame([
            metric_row("CT4", 100, month="2026-04-01", tactic="Search"),
            metric_row("CT4", 100, month="2026-05-01", tactic="Search"),
            metric_row("CT4", 100, month="2026-06-01", tactic="Search"),
        ])
        history = prepare_vehicle_historical_data(
            history_raw, slide_templates.tactic_config, vehicle_catalog,
            "Cadillac", ["CT4"], 50,
        )["CT4"]

        create_vehicle_insights_slide(
            presentation, "Cadillac", "Example Market", "June 2026",
            "CT4", active["CT4"], history, brand_configs,
            slide_templates, "ZONE", vehicle_catalog=vehicle_catalog,
        )
        create_vehicle_summary_slide(
            presentation, "Cadillac", "Example Market", "June 2026",
            "CT4", active["CT4"], brand_configs, slide_templates, "ZONE",
            vehicle_catalog=vehicle_catalog,
        )

        self.assertEqual(len(presentation.slides), 2)
        output = BytesIO()
        presentation.save(output)
        output.seek(0)
        reopened = Presentation(output)
        self.assertEqual(len(reopened.slides), 2)

    def test_missing_asset_renders_no_image_car(self):
        brand_configs = BrandConfigurations(BASE_PATH)
        slide_templates = SlideTemplates(brand_configs)
        presentation = Presentation()
        vehicle_catalog = get_vehicle_catalog(BASE_PATH)
        current = pd.DataFrame([metric_row("Escalade IQ", 100)])
        active, _ = prepare_vehicle_current_data(
            current, slide_templates.tactic_config, vehicle_catalog,
            "Cadillac", 50,
        )

        slide = create_vehicle_summary_slide(
            presentation, "Cadillac", "Example Market", "June 2026",
            "ESCALADE IQ", active["ESCALADE IQ"], brand_configs,
            slide_templates, "ZONE",
            vehicle_catalog=vehicle_catalog,
        )

        slide_text = "\n".join(
            shape.text for shape in slide.shapes if hasattr(shape, "text")
        )
        self.assertIn("NO IMAGE CAR", slide_text)


if __name__ == "__main__":
    unittest.main()
