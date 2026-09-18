import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import MagicMock, patch

import pandas as pd

from data_queries import DataQueries
from rollup_builder import (build_targets, catalog_query, generate_rollup_batch,
                            load_catalog, normalize_period, period_label, quarter_period, compact_deck_path,
                            report_filename_stem)


def sample_catalog():
    return pd.DataFrame([
        [brand, 'LMA', region, name, market, client]
        for brand in ('Buick', 'GMC', 'Cadillac', 'Chevrolet')
        for region, name, market, client in (
            ('SER', 'Panama City', 'PANFL', 'XTPC'),
            ('SER', 'Tallahassee', 'TALFL', 'XTPC'),
            ('NER', 'Albany', 'ALBGA', 'XALY'))
    ], columns=['Brand', 'Type', 'Region', 'MarketName', 'MarketCode', 'ClientCode'])


class RollupBuilderTests(unittest.TestCase):
    def test_market_suffix_uses_catalog_even_when_only_one_market_selected(self):
        targets, _ = build_targets(sample_catalog(), 'Cadillac', 'LMA', 'SER')
        self.assertEqual(report_filename_stem(targets[0], '2026-06-01', '2026-08-01'),
                         'XTPC-PANFL Cadillac June-Aug 2026 Summary')

    def test_requested_month_and_quarter_filenames(self):
        target = {'brand': 'Buick', 'client_code': 'XTPC', 'market_code': None, 'report_code': 'XTPC'}
        self.assertEqual(report_filename_stem(target, '2026-06-01', '2026-08-01'),
                         'XTPC Buick June-Aug 2026 Summary')
        self.assertEqual(report_filename_stem(target, '2026-07-01', '2026-09-01', 'Quarter'),
                         'XTPC Buick Q3 2026 Summary')
        self.assertEqual(report_filename_stem(target, '2025-12-01', '2026-02-01'),
                         'XTPC Buick Dec 2025-Feb 2026 Summary')
        with self.assertRaises(ValueError):
            report_filename_stem(target, '2026-06-01', '2026-08-01', 'Quarter')
        target.update(brand='GMC', market_code='PANFL', report_code='XTPC-PANFL')
        self.assertEqual(report_filename_stem(target, '2026-07-01', '2026-09-01', 'Quarter'),
                         'XTPC-PANFL GMC Q3 2026 Summary')

    def test_real_powerpoints_save_under_long_workspace_path(self):
        from pptx import Presentation
        from main_slide_generator import SlideGenerationOrchestrator
        generator = SlideGenerationOrchestrator.__new__(SlideGenerationOrchestrator)
        # Use the actual workspace root, not the short system temp directory.
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as temp:
            folder = Path(temp) / 'Rollup_2026-06_2026-08_20260918_122403_376827'
            for client, market in [('BBPG', 'BURVT'), ('XBAL', 'BALMD'), ('XBOS', 'BOSMA')]:
                prs = Presentation()
                prs.slides.add_slide(prs.slide_layouts[6])
                prefix = f'{client}-{market}_2026-06_2026-08'
                path = generator._save_complete_deck(
                    prs, 'Buick', None, '2026-08-01', 'LMA', market, 'NER', str(folder),
                    True, client_code=client, file_prefix=prefix, compact_output=True,
                    filename_stem=f'{client} Buick June-Aug 2026 Summary')
                self.assertLessEqual(len(str(Path(path).resolve())), 240)
                self.assertEqual(Path(path).parent, folder)
                self.assertEqual(Path(path).name, f'{client} Buick June-Aug 2026 Summary.pptx')
                self.assertEqual(len(Presentation(path).slides), 1)
                again = generator._save_complete_deck(
                    prs, 'Buick', None, '2026-08-01', 'LMA', market, 'NER', str(folder),
                    True, client_code=client, file_prefix=prefix, compact_output=True,
                    filename_stem=f'{client} Buick June-Aug 2026 Summary')
                self.assertNotEqual(path, again)
                self.assertEqual(len(Presentation(again).slides), 1)

    def test_compact_filename_keeps_identity_when_truncated(self):
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as folder:
            first = compact_deck_path(folder, 'Buick', 'LMA', 'NER', 'X' * 300 + 'A')
            second = compact_deck_path(folder, 'Buick', 'LMA', 'NER', 'X' * 300 + 'B')
            self.assertLessEqual(len(first), 240)
            self.assertNotEqual(first, second)

    def test_device_sign_in_callback_and_token_reuse(self):
        from powerbi_connector import PowerBIConnector
        connector = PowerBIConnector.__new__(PowerBIConnector)
        connector.access_token = None
        connector.token_expiry = 0
        connector.scope = ['test-scope']
        connector.app = MagicMock()
        connector.app.get_accounts.return_value = []
        connector.app.initiate_device_flow.return_value = {
            'user_code': 'TEST-CODE', 'verification_uri': 'https://microsoft.com/devicelogin'}
        connector.app.acquire_token_by_device_flow.return_value = {'access_token': 'test-token', 'expires_in': 3600}
        connector.session = MagicMock()
        connector.on_device_code = MagicMock()
        self.assertTrue(connector.authenticate())
        self.assertTrue(connector.authenticate())
        connector.on_device_code.assert_called_once_with('https://microsoft.com/devicelogin', 'TEST-CODE')
        connector.app.acquire_token_by_device_flow.assert_called_once()

    def test_quarters_and_cross_year_periods(self):
        for quarter in range(1, 5):
            start, end = quarter_period(2026, quarter)
            self.assertEqual(int(end[5:7]) - int(start[5:7]), 2)
        self.assertEqual(quarter_period(2026, 4), ('2026-10-01', '2026-12-01'))
        self.assertEqual(normalize_period('2025-12-29', '2026-02-28'), ('2025-12-01', '2026-02-01'))
        self.assertEqual(period_label('2025-12-01', '2026-02-01'), 'December 2025-February 2026')
        self.assertEqual(period_label('2026-06-01', '2026-06-01'), 'June 2026')
        with self.assertRaises(ValueError):
            normalize_period('2026-09-01', '2026-06-01')
        with self.assertRaises(ValueError):
            quarter_period(2026, 5)

    def test_catalog_query_includes_whole_period(self):
        query = catalog_query('GMC', 'LMA', '2026-04-01', '2026-06-01')
        self.assertIn('>= DATEVALUE("2026-04-01")', query)
        self.assertIn('<= DATEVALUE("2026-06-01")', query)
        self.assertIn('[Region]', query)
        self.assertIn('[Brand (Reporting)] = "GMC"', query)

    def test_all_brands_regions_and_xtpc_rules(self):
        catalog = sample_catalog()
        for brand in ('Buick', 'GMC', 'Cadillac', 'Chevrolet'):
            targets, issues = build_targets(catalog, brand, 'LMA', 'SER')
            self.assertFalse(issues)
            self.assertTrue(all(t['brand'] == brand and t['region'] == 'SER' for t in targets))
            self.assertEqual(len(targets), 1 if brand == 'Buick' else 2)
        buick, _ = build_targets(catalog, 'Buick', 'LMA', 'SER')
        self.assertIsNone(buick[0]['market_code'])
        gmc, _ = build_targets(catalog, 'GMC', 'LMA', 'SER')
        self.assertEqual({t['report_code'] for t in gmc}, {'XTPC-PANFL', 'XTPC-TALFL'})
        all_regions, _ = build_targets(catalog, 'GMC', 'LMA')
        self.assertEqual(len(all_regions), 3)

    def test_zone_aggregates_clients_and_reports_missing_mapping(self):
        catalog = sample_catalog().iloc[:2].copy()
        catalog['Type'] = 'ZONE'
        catalog['MarketCode'] = 'SAME'
        catalog['ClientCode'] = ['FIRST', 'SECOND']
        targets, issues = build_targets(catalog, 'Buick', 'ZONE')
        self.assertEqual(len(targets), 1)
        self.assertIsNone(targets[0]['client_code'])
        catalog.loc[0, 'MarketCode'] = ''
        _, issues = build_targets(catalog, 'Buick', 'ZONE')
        self.assertEqual(issues[0]['Rows'], 1)

    def test_region_is_applied_to_discovery_summary_and_ytd_queries(self):
        queries = DataQueries()
        for builder in (queries.get_tactic_site_combinations_query,
                        queries.get_executive_summary_query,
                        queries.get_ytd_kba_by_tactic_query,
                        queries.get_ytd_impressions_by_vehicle_query):
            for region in ('SER', ''):
                query = builder('GMC', None, '2026-08-01', zone_lma='LMA',
                                client_code='XTPC', market_code='PANFL', region=region)
                self.assertIn(f'COALESCE(\'PoP Master Table\'[Region], "") = "{region}"', query)
                self.assertIn('[Market Code] = "PANFL"', query)

    def test_load_catalog_handles_qualified_columns_and_blank_optional_fields(self):
        generator = MagicMock()
        generator.powerbi.execute_dax_query.return_value = pd.DataFrame([
            {'[Brand]': 'GMC', '[Type]': 'ZONE', '[MarketCode]': 'PANFL'}])
        catalog = load_catalog(generator, 'GMC', 'ZONE', '2026-06-01', '2026-08-01')
        self.assertEqual(catalog.ClientCode.tolist(), [''])
        self.assertEqual(catalog.Region.tolist(), [''])
        generator.powerbi.execute_dax_query.return_value = None
        with self.assertRaises(RuntimeError):
            load_catalog(generator, 'GMC', 'ZONE', '2026-06-01', '2026-08-01')

    def test_selected_batch_preserves_scope_and_persists_partial_failures(self):
        targets, _ = build_targets(sample_catalog(), 'GMC', 'LMA', 'SER')
        generator = MagicMock()
        generator.generate_quarter_rollup_for_market.side_effect = [
            {'filepath': 'first.pptx'}, RuntimeError('query failed')]
        with tempfile.TemporaryDirectory() as folder:
            result = generate_rollup_batch(generator, targets, '2025-12-01', '2026-02-01', output_folder=folder)
            saved = json.loads((Path(result['output_folder']) / 'rollup_results.json').read_text())
            self.assertEqual(saved['successful_decks'], 1)
            self.assertEqual(saved['errors'], 1)
            self.assertEqual(saved['ytd_start_month'], '2026-01-01')
            self.assertEqual(len(saved['details']), 2)
        for call in generator.generate_quarter_rollup_for_market.call_args_list:
            self.assertTrue(call.kwargs['compact_output'])
            self.assertTrue(call.kwargs['filename_stem'].endswith('GMC Dec 2025-Feb 2026 Summary'))
            self.assertEqual(call.kwargs['region'], 'SER')
            self.assertEqual(call.kwargs['client_code'], 'XTPC')
            self.assertEqual(call.kwargs['start_month_year'], '2025-12-01')


class RollupUiTests(unittest.TestCase):
    def test_selection_generation_and_stale_catalog_protection(self):
        from streamlit.testing.v1 import AppTest
        with patch('main_slide_generator.SlideGenerationOrchestrator') as cls, \
             patch('rollup_builder.load_catalog', return_value=sample_catalog()) as load, \
             patch('rollup_builder.generate_rollup_batch') as generate:
            app = AppTest.from_file(str(Path(__file__).with_name('rollup_ui.py')), default_timeout=30).run()
            self.assertFalse(app.exception)
            cls.assert_not_called()
            app.selectbox(key='rollup_brand').select('GMC').run()
            app.button(key='rollup_load').click().run()
            app.selectbox(key='rollup_region').select('SER').run()
            app.button(key='rollup_select_all').click().run()
            self.assertEqual(len(app.multiselect(key='rollup_markets').value), 2)
            self.assertFalse(app.button(key='rollup_generate').disabled)
            generate.return_value = {'start_month': '2026-06-01', 'end_month': '2026-08-01',
                                     'successful_decks': 2, 'expected_decks': 2, 'errors': 0,
                                     'output_folder': 'test-output', 'details': [{'Status': 'Success'}]}
            app.button(key='rollup_generate').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(len(generate.call_args.args[1]), 2)
            self.assertEqual(generate.call_args.kwargs['period_mode'], 'Month range')
            self.assertEqual(len(app.success), 1)
            app.selectbox(key='rollup_region').select('NER').run()
            self.assertEqual(app.multiselect(key='rollup_markets').value, [])
            app.selectbox(key='rollup_brand').select('Buick').run()
            self.assertFalse(any(b.key == 'rollup_generate' for b in app.button))
            self.assertEqual(load.call_count, 1)
            app.selectbox(key='rollup_period_mode').select('Quarter').run()
            app.selectbox(key='rollup_quarter').select(1).run()
            app.button(key='rollup_load').click().run()
            self.assertTrue(load.call_args.args[-2].endswith('-01-01'))
            self.assertTrue(load.call_args.args[-1].endswith('-03-01'))
            app.selectbox(key='rollup_region').select('SER').run()
            app.button(key='rollup_select_all').click().run()
            self.assertEqual(len(app.multiselect(key='rollup_markets').value), 1)
            app.button(key='rollup_clear').click().run()
            self.assertEqual(app.multiselect(key='rollup_markets').value, [])
            self.assertFalse(app.exception)


if __name__ == '__main__':
    unittest.main()
