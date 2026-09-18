"""Vehicle-level Insights and Summary slides.

The module keeps data preparation separate from rendering so the spend rule and
asset matching can be tested without creating a PowerPoint presentation.
"""

import os
from io import BytesIO

import matplotlib.pyplot as plt
import pandas as pd
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

import settings
from utils import find_column, get_safe_font, register_brand_fonts, rgb_to_matplotlib
from vehicle_catalog import get_vehicle_catalog


METRIC_COLUMNS = (
    "Total Cost",
    "Impressions",
    "Video Completions",
    "Video Plays",
    "Clicks",
    "Total Conversions",
)

def _normalize_site(site, site_grouping_rules):
    site_text = "" if site is None else str(site).strip()
    site_upper = site_text.upper()

    for group_name, site_list in (site_grouping_rules or {}).items():
        group_upper = str(group_name).upper().strip()
        normalized_sites = {str(item).upper().strip() for item in site_list}
        if site_upper in normalized_sites or (group_upper and group_upper in site_upper):
            return str(group_name)

    if "MSP" in site_upper:
        if "AMPERSAND" in site_upper:
            return "Ampersand"
        return f"MSP {site_text}"
    return site_text


def _standardize_vehicle_metrics(raw_df, tactic_config, vehicle_catalog, brand,
                                 include_month=False):
    if raw_df is None or raw_df.empty:
        return pd.DataFrame()

    vehicle_col = find_column(raw_df, "Vehicle")
    if not vehicle_col:
        return pd.DataFrame()

    column_names = {
        "Tactic (Reporting)": find_column(raw_df, "Tactic (Reporting)"),
        "Strategy (groups)": find_column(raw_df, "Strategy (groups)"),
        "Site (Reporting)": find_column(raw_df, "Site (Reporting)"),
        "Audience": find_column(raw_df, "Audience"),
    }
    if include_month:
        column_names["Month Year"] = find_column(raw_df, "Month Year")

    work = pd.DataFrame(index=raw_df.index)
    work["Source Vehicle"] = raw_df[vehicle_col].fillna("").astype(str).str.strip()
    work = work[work["Source Vehicle"] != ""].copy()
    if work.empty:
        return work

    resolutions = work["Source Vehicle"].map(
        lambda vehicle: vehicle_catalog.resolve(brand, vehicle)
    )
    work["Vehicle"] = resolutions.map(lambda item: item.canonical_name)
    work["Vehicle Key"] = resolutions.map(lambda item: item.canonical_key)
    work["Vehicle Classification"] = resolutions.map(
        lambda item: item.classification
    )
    work["Vehicle Status"] = resolutions.map(lambda item: item.status)
    work["Generate Vehicle Slides"] = resolutions.map(
        lambda item: item.generate_slides
    )
    work["Vehicle Taxonomy Reason"] = resolutions.map(lambda item: item.reason)
    work["Vehicle Taxonomy Mapped"] = resolutions.map(lambda item: item.mapped)

    for output_name, source_name in column_names.items():
        if source_name:
            work[output_name] = raw_df[source_name]
        else:
            work[output_name] = ""

    for text_column in ("Tactic (Reporting)", "Strategy (groups)",
                        "Site (Reporting)", "Audience"):
        work[text_column] = work[text_column].fillna("").astype(str).str.strip()

    work.loc[work["Tactic (Reporting)"] == "", "Tactic (Reporting)"] = "Unknown"

    for metric_name in METRIC_COLUMNS:
        metric_col = find_column(raw_df, metric_name)
        work[metric_name] = (
            pd.to_numeric(raw_df[metric_col], errors="coerce").fillna(0)
            if metric_col else 0
        )

    if include_month:
        work["Month Year"] = pd.to_datetime(work["Month Year"], errors="coerce")
        work = work[work["Month Year"].notna()].copy()

    work["Effective Tactic"] = work.apply(
        lambda row: tactic_config._resolve_display_tactic(
            row["Tactic (Reporting)"],
            row["Strategy (groups)"],
            row["Site (Reporting)"],
        ),
        axis=1,
    )
    site_rules = tactic_config.rules.get("site_grouping_rules", {})
    work["Grouped Site"] = work["Site (Reporting)"].map(
        lambda site: _normalize_site(site, site_rules)
    )
    return work


def _aggregate_spend_grain(work, include_month=False):
    if work is None or work.empty:
        return pd.DataFrame()

    group_columns = ["Vehicle Key"]
    if include_month:
        group_columns.append("Month Year")
    group_columns.extend(["Effective Tactic", "Grouped Site", "Audience"])

    aggregations = {
        "Vehicle": "first",
        "Source Vehicle": lambda values: ", ".join(
            sorted({str(value).strip() for value in values if str(value).strip()})
        ),
        "Vehicle Classification": "first",
        "Vehicle Status": "first",
        "Generate Vehicle Slides": "first",
        "Vehicle Taxonomy Reason": "first",
        "Vehicle Taxonomy Mapped": "first",
    }
    aggregations.update({metric_name: "sum" for metric_name in METRIC_COLUMNS})
    return work.groupby(
        group_columns, as_index=False, dropna=False, sort=False
    ).agg(aggregations)


def _presentation_columns(frame):
    result = frame.rename(columns={
        "Effective Tactic": "Tactic (Reporting)",
        "Grouped Site": "Site (Reporting)",
    }).copy()
    return result.drop(columns=[
        "Vehicle Key",
        "Source Vehicle",
        "Vehicle Classification",
        "Vehicle Status",
        "Generate Vehicle Slides",
        "Vehicle Taxonomy Reason",
        "Vehicle Taxonomy Mapped",
    ], errors="ignore")


def prepare_vehicle_current_data(raw_df, tactic_config, vehicle_catalog, brand,
                                 minimum_spend):
    """Apply the inclusive current-month spend rule and identify omitted models.

    Returns:
        (active_models, excluded_models), where active_models maps each approved
        canonical JSON vehicle to qualifying grouped rows. Exclusions include
        both spend failures and JSON taxonomy decisions.
    """
    threshold = max(float(minimum_spend or 0), 0.0)
    standardized = _standardize_vehicle_metrics(
        raw_df, tactic_config, vehicle_catalog, brand
    )
    grouped = _aggregate_spend_grain(standardized)
    if grouped.empty:
        return {}, []

    active_models = {}
    excluded_models = []
    for _, vehicle_rows in grouped.groupby("Vehicle Key", sort=True):
        vehicle_name = str(vehicle_rows["Vehicle"].iloc[0]).strip()
        spend = pd.to_numeric(vehicle_rows["Total Cost"], errors="coerce").fillna(0)
        qualifying_mask = spend >= threshold
        generate_slides = bool(vehicle_rows["Generate Vehicle Slides"].iloc[0])

        if not generate_slides:
            excluded_models.append({
                "vehicle": vehicle_name,
                "source_values": str(vehicle_rows["Source Vehicle"].iloc[0]),
                "exclusion_type": "taxonomy",
                "classification": str(
                    vehicle_rows["Vehicle Classification"].iloc[0]
                ),
                "status": str(vehicle_rows["Vehicle Status"].iloc[0]),
                "reason": str(
                    vehicle_rows["Vehicle Taxonomy Reason"].iloc[0]
                ),
                "total_spend": float(spend.sum()),
                "max_combination_spend": (
                    float(spend.max()) if not spend.empty else 0.0
                ),
                "combination_count": int(len(vehicle_rows)),
                "threshold": threshold,
            })
            continue

        if qualifying_mask.any():
            active_models[vehicle_name] = _presentation_columns(
                vehicle_rows.loc[qualifying_mask].reset_index(drop=True)
            )
        else:
            excluded_models.append({
                "vehicle": vehicle_name,
                "source_values": str(vehicle_rows["Source Vehicle"].iloc[0]),
                "exclusion_type": "spend",
                "classification": str(
                    vehicle_rows["Vehicle Classification"].iloc[0]
                ),
                "status": str(vehicle_rows["Vehicle Status"].iloc[0]),
                "reason": "",
                "total_spend": float(spend.sum()),
                "max_combination_spend": float(spend.max()) if not spend.empty else 0.0,
                "combination_count": int(len(vehicle_rows)),
                "threshold": threshold,
            })

    return dict(sorted(active_models.items(), key=lambda item: item[0].lower())), excluded_models


def prepare_vehicle_historical_data(raw_df, tactic_config, vehicle_catalog, brand,
                                    active_vehicle_names, minimum_spend,
                                    apply_spend_filter=True):
    """Prepare three-month history using the same spend grain as current data."""
    if not active_vehicle_names:
        return {}

    threshold = max(float(minimum_spend or 0), 0.0)
    standardized = _standardize_vehicle_metrics(
        raw_df, tactic_config, vehicle_catalog, brand, include_month=True
    )
    grouped = _aggregate_spend_grain(standardized, include_month=True)
    if grouped.empty:
        return {vehicle: pd.DataFrame() for vehicle in active_vehicle_names}

    active_lookup = {}
    for vehicle in active_vehicle_names:
        resolution = vehicle_catalog.resolve(brand, vehicle)
        active_lookup[resolution.canonical_key] = resolution.canonical_name
    grouped = grouped[grouped["Vehicle Key"].isin(active_lookup)].copy()
    if apply_spend_filter:
        spend = pd.to_numeric(grouped["Total Cost"], errors="coerce").fillna(0)
        grouped = grouped[spend >= threshold].copy()

    result = {}
    for vehicle_key, display_name in active_lookup.items():
        vehicle_rows = grouped[grouped["Vehicle Key"] == vehicle_key].copy()
        if vehicle_rows.empty:
            result[display_name] = pd.DataFrame()
            continue

        aggregations = {metric_name: "sum" for metric_name in METRIC_COLUMNS}
        history = vehicle_rows.groupby(
            ["Month Year", "Effective Tactic"],
            as_index=False,
            sort=True,
        ).agg(aggregations)
        history = history.rename(columns={"Effective Tactic": "Tactic (Reporting)"})
        result[display_name] = history.sort_values("Month Year").reset_index(drop=True)
    return result


def resolve_vehicle_image_path(vehicle_name, brand, slide_kind, base_path,
                               vehicle_catalog=None):
    """Resolve an image exclusively through ``vehicle_images.json``."""
    catalog = vehicle_catalog or get_vehicle_catalog(base_path)
    return catalog.resolve_image_path(brand, vehicle_name, slide_kind)


def _get_tactic_color_map(brand):
    raw_colors = settings.get_tactic_colors(brand)
    tactic_colors = {}
    for tactic, rgb_values in raw_colors.items():
        if isinstance(rgb_values, list) and len(rgb_values) == 3:
            tactic_colors[str(tactic).lower().strip()] = RGBColor(*rgb_values)

    defaults = {
        "fep": RGBColor(0, 119, 217),
        "preroll": RGBColor(234, 33, 43),
        "youtube": RGBColor(179, 179, 179),
        "search": RGBColor(242, 188, 24),
        "display": RGBColor(242, 188, 24),
        "social video": RGBColor(51, 173, 255),
        "audio": RGBColor(51, 51, 51),
    }
    for tactic, color in defaults.items():
        tactic_colors.setdefault(tactic, color)
    return tactic_colors


def _get_tactic_rgb(tactic_name, tactic_color_map):
    tactic_key = str(tactic_name or "").lower().strip()
    return tactic_color_map.get(tactic_key, RGBColor(0, 119, 217))


def _add_fit_vehicle_image(slide, image_path, left, top, max_width, max_height,
                           brand_config):
    if image_path and os.path.isfile(image_path):
        try:
            with Image.open(image_path) as image:
                original_width, original_height = image.size
            if original_width > 0 and original_height > 0:
                scale = min(
                    max_width / float(original_width),
                    max_height / float(original_height),
                )
                width = original_width * scale
                height = original_height * scale
                slide.shapes.add_picture(
                    image_path,
                    Inches(left + (max_width - width) / 2),
                    Inches(top + (max_height - height) / 2),
                    width=Inches(width),
                    height=Inches(height),
                )
                return True
        except Exception as exc:
            print(f"  [WARN] Vehicle image error for '{image_path}': {exc}")

    placeholder = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(max_width), Inches(max_height)
    )
    placeholder.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = placeholder.text_frame.paragraphs[0]
    paragraph.text = "NO IMAGE CAR"
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.name = brand_config.get("primary_font", "Segoe UI")
    paragraph.font.size = Pt(20)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(190, 55, 55)
    return False


def _add_common_slide_structure(slide, brand_config, slide_templates,
                                market_name, vehicle_name, slide_label,
                                formatted_month_year, zone_lma_type):
    layouts = slide_templates.layouts
    px = slide_templates.pixels_to_inches
    layouts.add_data_background(slide, brand_config)
    layouts.add_data_brand_logo(slide, brand_config)
    layouts.add_data_logo_divider(slide, brand_config)
    layouts.add_mrg_logo(slide, position="data")
    layouts.add_gm_confidential(slide, brand_config, position="data")

    font_name = brand_config.get("primary_font", "Segoe UI")
    text_color = brand_config["text_color_dark"]

    market_box = slide.shapes.add_textbox(px(0), px(25), px(1280), px(40))
    market_paragraph = market_box.text_frame.paragraphs[0]
    market_paragraph.text = str(market_name).upper()
    market_paragraph.alignment = PP_ALIGN.CENTER
    market_paragraph.font.name = font_name
    market_paragraph.font.size = Pt(22)
    market_paragraph.font.bold = True
    market_paragraph.font.color.rgb = text_color

    divider = slide.shapes.add_connector(1, px(485), px(70), px(799), px(70))
    divider.line.color.rgb = brand_config["line_color"]
    divider.line.width = Pt(brand_config.get("line_width", 2))
    divider.shadow.inherit = False

    title_box = slide.shapes.add_textbox(px(0), px(75), px(1280), px(45))
    title_paragraph = title_box.text_frame.paragraphs[0]
    title_paragraph.text = f"{str(vehicle_name).upper().strip()} {slide_label}"
    title_paragraph.alignment = PP_ALIGN.CENTER
    title_paragraph.font.name = font_name
    title_paragraph.font.size = Pt(32)
    title_paragraph.font.bold = True
    title_paragraph.font.color.rgb = text_color

    month_box = slide.shapes.add_textbox(px(0), px(122), px(1280), px(30))
    month_paragraph = month_box.text_frame.paragraphs[0]
    month_paragraph.text = formatted_month_year
    month_paragraph.alignment = PP_ALIGN.CENTER
    month_paragraph.font.name = font_name
    month_paragraph.font.size = Pt(14)
    month_paragraph.font.bold = True
    month_paragraph.font.color.rgb = text_color

    layouts.add_source_label(
        slide, brand_config, zone_lma_type, formatted_month_year
    )


def _metric_rows_by_tactic(current_df):
    if current_df is None or current_df.empty:
        return [], []
    tactic_col = find_column(current_df, "Tactic (Reporting)")
    if not tactic_col:
        return [], []

    calculation = current_df.copy()
    resolved_columns = {}
    for metric_name in METRIC_COLUMNS:
        metric_col = find_column(calculation, metric_name)
        resolved_columns[metric_name] = metric_col
        if metric_col:
            calculation[metric_col] = pd.to_numeric(
                calculation[metric_col], errors="coerce"
            ).fillna(0)

    grouped = calculation.groupby(tactic_col, as_index=False).sum(numeric_only=True)
    vcr_rows = []
    cpo_rows = []
    for _, row in grouped.iterrows():
        tactic = row[tactic_col]
        completions = row.get(resolved_columns["Video Completions"], 0)
        plays = row.get(resolved_columns["Video Plays"], 0)
        cost = row.get(resolved_columns["Total Cost"], 0)
        conversions = row.get(resolved_columns["Total Conversions"], 0)
        if plays > 0:
            vcr_rows.append((tactic, f"{completions / plays:.0%}"))
        if conversions > 0:
            cpo_rows.append((tactic, f"${cost / conversions:,.2f}"))
    return vcr_rows, cpo_rows


def _add_metric_table(slide, metric_label, rows, left, top, width, font_name):
    if not rows:
        return
    row_count = len(rows) + 1
    height = min(2.42, max(0.62, row_count * 0.29))
    table = slide.shapes.add_table(
        row_count, 2, Inches(left), Inches(top), Inches(width), Inches(height)
    ).table
    table.columns[0].width = Inches(width * 0.64)
    table.columns[1].width = Inches(width * 0.36)

    for column_index, label in enumerate(("Tactic", metric_label)):
        cell = table.cell(0, column_index)
        cell.text = label
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
        paragraph = cell.text_frame.paragraphs[0]
        paragraph.font.name = font_name
        paragraph.font.size = Pt(12)
        paragraph.font.bold = True
        paragraph.font.color.rgb = RGBColor(0, 0, 0)
        paragraph.alignment = PP_ALIGN.RIGHT if column_index else PP_ALIGN.LEFT

    font_size = 10 if len(rows) > 7 else 11
    for row_index, (tactic, value) in enumerate(rows, start=1):
        fill_color = RGBColor(245, 245, 245) if row_index % 2 == 0 else RGBColor(255, 255, 255)
        for column_index, text_value in enumerate((str(tactic), value)):
            cell = table.cell(row_index, column_index)
            cell.text = text_value
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill_color
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.font.name = font_name
            paragraph.font.size = Pt(font_size)
            paragraph.font.color.rgb = RGBColor(35, 35, 35)
            paragraph.alignment = PP_ALIGN.RIGHT if column_index else PP_ALIGN.LEFT


def _add_no_data_label(slide, text, left, top, width, height, brand_config):
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.name = brand_config.get("primary_font", "Segoe UI")
    paragraph.font.size = Pt(12)
    paragraph.font.color.rgb = brand_config["text_color_dark"]


def _chart_style(brand_config):
    is_dark = brand_config.get("name", "").lower() == "cadillac"
    return {
        "foreground": "#FFFFFF" if is_dark else "#222222",
        "grid": "#777777" if is_dark else "#D0D0D0",
        "spine": "#AAAAAA" if is_dark else "#CCCCCC",
    }


def _add_insights_line_charts(slide, historical_df, tactic_color_map,
                              brand_config, base_path):
    if historical_df is None or historical_df.empty:
        _add_no_data_label(
            slide, "NO QUALIFYING 3-MONTH HISTORY", 0.75, 4.65, 11.8, 1.6,
            brand_config,
        )
        return

    date_col = find_column(historical_df, "Month Year")
    tactic_col = find_column(historical_df, "Tactic (Reporting)")
    impressions_col = find_column(historical_df, "Impressions")
    conversions_col = find_column(historical_df, "Total Conversions")
    if not date_col or not tactic_col or not impressions_col or not conversions_col:
        _add_no_data_label(
            slide, "NO QUALIFYING 3-MONTH HISTORY", 0.75, 4.65, 11.8, 1.6,
            brand_config,
        )
        return

    history = historical_df.copy()
    history[date_col] = pd.to_datetime(history[date_col], errors="coerce")
    history = history[history[date_col].notna()].sort_values(date_col)
    history[impressions_col] = pd.to_numeric(
        history[impressions_col], errors="coerce"
    ).fillna(0)
    history[conversions_col] = pd.to_numeric(
        history[conversions_col], errors="coerce"
    ).fillna(0)
    history["Month Label"] = history[date_col].dt.strftime("%b-%y")
    month_labels = list(
        history[[date_col, "Month Label"]]
        .drop_duplicates()
        .sort_values(date_col)["Month Label"]
    )

    register_brand_fonts(base_path)
    font_family, _ = get_safe_font(brand_config.get("primary_font", "Arial"))
    style = _chart_style(brand_config)

    for metric_col, title, left in (
        (impressions_col, "Impressions by Tactic", 0.75),
        (conversions_col, "Outcomes by Tactic", 6.75),
    ):
        pivot = history.pivot_table(
            index="Month Label",
            columns=tactic_col,
            values=metric_col,
            aggfunc="sum",
            fill_value=0,
        ).reindex(month_labels).fillna(0)

        figure, axis = plt.subplots(figsize=(5.5, 2.45), dpi=150)
        figure.patch.set_alpha(0)
        axis.patch.set_alpha(0)
        for tactic in pivot.columns:
            color = rgb_to_matplotlib(_get_tactic_rgb(tactic, tactic_color_map))
            axis.plot(
                pivot.index, pivot[tactic], marker="o", label=tactic,
                color=color, linewidth=2,
            )
        axis.set_title(
            title, fontsize=11, fontweight="bold", pad=8,
            color=style["foreground"], fontfamily=font_family,
        )
        axis.grid(True, linestyle="--", alpha=0.45, axis="y", color=style["grid"])
        axis.spines["top"].set_visible(False)
        axis.spines["right"].set_visible(False)
        axis.spines["left"].set_visible(False)
        axis.spines["bottom"].set_color(style["spine"])
        axis.yaxis.set_major_formatter(
            plt.FuncFormatter(lambda value, _: f"{int(value):,}")
        )
        axis.tick_params(axis="both", labelsize=8, colors=style["foreground"])
        legend = axis.legend(
            loc="upper center", bbox_to_anchor=(0.5, -0.15), ncol=4,
            frameon=False, fontsize=8,
        )
        if legend:
            for legend_text in legend.get_texts():
                legend_text.set_color(style["foreground"])
                legend_text.set_fontfamily(font_family)
        figure.tight_layout()

        buffer = BytesIO()
        figure.savefig(buffer, format="png", transparent=True)
        plt.close(figure)
        buffer.seek(0)
        slide.shapes.add_picture(
            buffer, Inches(left), Inches(4.48), width=Inches(5.75)
        )


def create_vehicle_insights_slide(prs, brand, market_name,
                                  formatted_month_year, vehicle_name,
                                  current_df, historical_df, brand_configs,
                                  slide_templates, zone_lma_type,
                                  vehicle_catalog=None):
    brand_config = brand_configs.get_brand_config(brand)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_common_slide_structure(
        slide, brand_config, slide_templates, market_name, vehicle_name,
        "INSIGHTS", formatted_month_year, zone_lma_type,
    )

    image_path = resolve_vehicle_image_path(
        vehicle_name, brand, "insights", brand_configs.base_path,
        vehicle_catalog=vehicle_catalog,
    )
    _add_fit_vehicle_image(
        slide, image_path, 0.72, 1.55, 5.35, 2.65, brand_config
    )

    vcr_rows, cpo_rows = _metric_rows_by_tactic(current_df)
    font_name = brand_config.get("primary_font", "Segoe UI")
    if vcr_rows and cpo_rows:
        _add_metric_table(slide, "VCR", vcr_rows, 7.25, 1.62, 2.55, font_name)
        _add_metric_table(slide, "CPO", cpo_rows, 10.05, 1.62, 2.55, font_name)
    elif vcr_rows:
        _add_metric_table(slide, "VCR", vcr_rows, 8.05, 1.62, 4.1, font_name)
    elif cpo_rows:
        _add_metric_table(slide, "CPO", cpo_rows, 8.05, 1.62, 4.1, font_name)
    else:
        _add_no_data_label(
            slide, "NO VCR OR CPO DATA", 7.25, 1.75, 5.1, 1.5, brand_config
        )

    _add_insights_line_charts(
        slide, historical_df, _get_tactic_color_map(brand),
        brand_config, brand_configs.base_path,
    )
    return slide


def _sum_metric(frame, metric_name):
    metric_col = find_column(frame, metric_name)
    if not metric_col:
        return 0.0
    return float(pd.to_numeric(frame[metric_col], errors="coerce").fillna(0).sum())


def _add_summary_cards(slide, current_df, brand_config):
    total_impressions = _sum_metric(current_df, "Impressions")
    total_completions = _sum_metric(current_df, "Video Completions")
    total_plays = _sum_metric(current_df, "Video Plays")
    total_conversions = _sum_metric(current_df, "Total Conversions")
    total_cost = _sum_metric(current_df, "Total Cost")
    vcr = total_completions / total_plays if total_plays > 0 else 0
    cpo = total_cost / total_conversions if total_conversions > 0 else 0

    cards = (
        ("Impressions", f"{int(total_impressions):,}"),
        ("Video Completes", f"{int(total_completions):,}"),
        ("Total Conversions", f"{int(total_conversions):,}"),
        ("VCR", f"{vcr:.0%}"),
        ("CPO", f"${cpo:,.2f}"),
    )
    metric_box = brand_config.get("metric_box", {})
    fill_color = metric_box.get("fill_color", RGBColor(255, 255, 255))
    is_dark = brand_config.get("name", "").lower() == "cadillac"
    label_color = RGBColor(215, 215, 215) if is_dark else RGBColor(105, 105, 105)
    value_color = RGBColor(255, 255, 255) if is_dark else RGBColor(40, 40, 40)
    font_name = brand_config.get("primary_font", "Segoe UI")

    for index, (label, value) in enumerate(cards):
        left = 0.62 + index * 2.52
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(left), Inches(1.67), Inches(2.30), Inches(0.86),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
        if metric_box.get("has_border", False) and metric_box.get("border_color"):
            shape.line.color.rgb = metric_box["border_color"]
            shape.line.width = Pt(metric_box.get("border_width", 1))
        else:
            shape.line.fill.background()

        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_top = Inches(0.08)
        frame.margin_bottom = Inches(0.03)
        label_paragraph = frame.paragraphs[0]
        label_paragraph.text = label
        label_paragraph.alignment = PP_ALIGN.CENTER
        label_paragraph.font.name = font_name
        label_paragraph.font.size = Pt(10)
        label_paragraph.font.color.rgb = label_color
        value_paragraph = frame.add_paragraph()
        value_paragraph.text = value
        value_paragraph.alignment = PP_ALIGN.CENTER
        value_paragraph.font.name = font_name
        value_paragraph.font.size = Pt(22)
        value_paragraph.font.bold = True
        value_paragraph.font.color.rgb = value_color


def _add_spend_donut_chart(slide, current_df, tactic_color_map,
                           brand_config, base_path):
    tactic_col = find_column(current_df, "Tactic (Reporting)")
    cost_col = find_column(current_df, "Total Cost")
    if current_df is None or current_df.empty or not tactic_col or not cost_col:
        _add_no_data_label(
            slide, "NO DELIVERED SPEND", 6.8, 3.15, 5.6, 2.0, brand_config
        )
        return

    spend = current_df[[tactic_col, cost_col]].copy()
    spend[cost_col] = pd.to_numeric(spend[cost_col], errors="coerce").fillna(0)
    spend = spend.groupby(tactic_col, as_index=False)[cost_col].sum()
    spend = spend[spend[cost_col] > 0]
    if spend.empty:
        _add_no_data_label(
            slide, "NO DELIVERED SPEND", 6.8, 3.15, 5.6, 2.0, brand_config
        )
        return

    register_brand_fonts(base_path)
    font_family, _ = get_safe_font(brand_config.get("primary_font", "Arial"))
    style = _chart_style(brand_config)
    tactics = spend[tactic_col].tolist()
    costs = spend[cost_col].tolist()
    colors = [
        rgb_to_matplotlib(_get_tactic_rgb(tactic, tactic_color_map))
        for tactic in tactics
    ]

    figure, axis = plt.subplots(figsize=(5.2, 3.2), dpi=150)
    figure.patch.set_alpha(0)
    axis.patch.set_alpha(0)
    wedges, _ = axis.pie(
        costs, labels=None, colors=colors, startangle=90,
        wedgeprops={"width": 0.4, "edgecolor": "white", "linewidth": 2},
    )
    axis.set_title(
        "Delivered Spend by Tactic", fontsize=12, fontweight="bold", pad=10,
        color=style["foreground"], fontfamily=font_family,
    )
    legend = axis.legend(
        wedges, tactics, loc="center left", bbox_to_anchor=(1, 0.5),
        frameon=False, fontsize=9,
    )
    for legend_text in legend.get_texts():
        legend_text.set_color(style["foreground"])
        legend_text.set_fontfamily(font_family)
    figure.tight_layout()

    buffer = BytesIO()
    figure.savefig(buffer, format="png", transparent=True)
    plt.close(figure)
    buffer.seek(0)
    slide.shapes.add_picture(
        buffer, Inches(6.65), Inches(2.75), width=Inches(5.75)
    )


def create_vehicle_summary_slide(prs, brand, market_name,
                                 formatted_month_year, vehicle_name,
                                 current_df, brand_configs, slide_templates,
                                 zone_lma_type, vehicle_catalog=None):
    brand_config = brand_configs.get_brand_config(brand)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_common_slide_structure(
        slide, brand_config, slide_templates, market_name, vehicle_name,
        "SUMMARY", formatted_month_year, zone_lma_type,
    )
    _add_summary_cards(slide, current_df, brand_config)

    image_path = resolve_vehicle_image_path(
        vehicle_name, brand, "exec", brand_configs.base_path,
        vehicle_catalog=vehicle_catalog,
    )
    _add_fit_vehicle_image(
        slide, image_path, 0.72, 2.82, 5.35, 3.15, brand_config
    )
    _add_spend_donut_chart(
        slide, current_df, _get_tactic_color_map(brand),
        brand_config, brand_configs.base_path,
    )
    return slide
